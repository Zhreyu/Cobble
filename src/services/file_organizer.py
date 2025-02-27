#!/usr/bin/env python3

import os
import sys
import json
import re
import shutil
import logging
from datetime import datetime
from typing import Dict, List, Optional, Tuple

import PyPDF2
import docx
import openpyxl
from tqdm import tqdm
from groq import Groq

from src.context.file_organizer_context import FileOrganizerContext
from src.models.rule import Rule

logger = logging.getLogger(__name__)

class SmartFileSorter:
    def __init__(self, api_key: str, model_name: str = "mixtral-8x7b-32768"):
        self.context = FileOrganizerContext()
        self.client = Groq(api_key=api_key)
        self.model_name = model_name
        # Keep track of original file layout (relative paths).
        self.original_files: Dict[str, List[str]] = {}

    def parse_excel_file(self, file_path: str) -> str:
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            sheet = workbook.worksheets[0]
            rows_data = []
            for row in sheet.iter_rows(values_only=True):
                if any(row):
                    row_text = " ".join(str(cell) for cell in row if cell is not None)
                    rows_data.append(row_text)
            return "\n".join(rows_data[:50])  # first 50 rows
        except Exception as e:
            logger.error(f"Error parsing Excel file {file_path}: {str(e)}")
            return ""

    def get_file_excerpt(self, file_path: str, max_words: int = 500) -> str:
        """
        Return a short text excerpt from the file (similar to get_file_content, 
        but we won't do per-file classification here).
        """
        ext = os.path.splitext(file_path)[1].lower()
        text = ""
        try:
            if ext in ('.txt', '.md'):
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    text = f.read(max_words * 5)
            elif ext == '.pdf':
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages[:1]:  # just 1 page
                        text += page.extract_text()
            elif ext in ('.xlsx', '.xls'):
                text = self.parse_excel_file(file_path)
            elif ext in ('.docx', '.doc'):
                doc = docx.Document(file_path)
                text = "\n".join(p.text for p in doc.paragraphs[:10])

        except Exception as e:
            logger.error(f"Error reading file {file_path}: {str(e)}")
            text = os.path.basename(file_path)  # fallback

        # Truncate if large
        words = text.split()
        if len(words) > max_words:
            text = " ".join(words[:max_words])

        return text

    def bulk_classify_all_files(self, directory: str) -> Dict[str, str]:
        """
        Collect an excerpt from every file, send them ALL to the LLM in ONE request,
        parse a JSON mapping of {filename: Category}.
        
        Returns: { "sample_1.txt": "Business/Startup", "sample_2.pdf": "Academic", ... }
        """
        # 1) Build a list of {filename, extension, excerpt}
        file_info_list = []
        for root, _, files in os.walk(directory):
            for f in files:
                file_path = os.path.join(root, f)
                if not os.path.isfile(file_path):
                    continue
                ext = os.path.splitext(f)[1]
                excerpt = self.get_file_excerpt(file_path, max_words=300)
                file_info_list.append({
                    "filename": f,
                    "extension": ext,
                    "excerpt": excerpt
                })

        # 2) Build the prompt for a single bulk classification
        system_instructions = """You are a file categorization assistant. 
We have multiple files, each with:
- filename
- extension
- a short content excerpt

Your job:
1) Assign each file EXACTLY ONE category from:
   - Academic
   - Business
   - Documents
   - Personal
   - Media
2) Optionally use ONE subcategory, e.g. "Academic/Physics".
3) Return a single JSON object with the structure:
{
   "filename1.ext": "MainCategory[/SubCategory]",
   "filename2.ext": "MainCategory[/SubCategory]",
   ...
}
Follow these rules:
- Use PascalCase for categories/subcategories, e.g. "Academic/Mathematics"
- Return ONLY valid JSON, with no extra text or comments.
- If unsure, guess the most likely category from the excerpt or extension.
"""

        user_content_list = []
        for item in file_info_list:
            # We'll limit how much excerpt we pass
            excerpt = item['excerpt'][:600]  # extra guard
            msg = (
                f"Filename: {item['filename']}\n"
                f"Extension: {item['extension']}\n"
                f"Excerpt: {excerpt}\n---\n"
            )
            user_content_list.append(msg)

        # Combine the user message
        user_message = "\n".join(user_content_list)

        messages = [
            {"role": "system", "content": system_instructions},
            {"role": "user", "content": user_message}
        ]

        # 3) Call the LLM once
        try:
            response = self.client.chat.completions.create(
                model=self.model_name,
                messages=messages,
                temperature=0.0,   # Lower temp for more consistent responses
                max_tokens=800
            )
            raw_text = response.choices[0].message.content.strip()
            # Try to parse JSON
            json_text = self.extract_json_from_text(raw_text)
            mapping = json.loads(json_text)  # { filename -> "Category/Sub" }

            # Validate each category
            final_mapping = {}
            for fname, cat in mapping.items():
                cat_str = cat.strip().replace('"', '').replace("'", "")
                cat_str = cat_str.split('.')[0]
                cat_str = re.sub(r'[^\w/]', '', cat_str)
                if not re.match(r'^[A-Z][a-zA-Z]+(/[A-Z][a-zA-Z]+)?$', cat_str):
                    cat_str = "Documents/Uncategorized"
                final_mapping[fname] = cat_str

            return final_mapping

        except Exception as e:
            logger.error(f"Error in bulk classification: {str(e)}")
            # fallback: everything -> "Documents/Uncategorized"
            fallback_map = {}
            for item in file_info_list:
                fallback_map[item['filename']] = "Documents/Uncategorized"
            return fallback_map

    def extract_json_from_text(self, text: str) -> str:
        """
        A naive helper to extract a JSON block from text.
        We find first '{' and last '}' and slice. 
        If not found, return "{}".
        """
        start = text.find('{')
        end = text.rfind('}')
        if start == -1 or end == -1:
            return "{}"
        return text[start:end+1]

    def preview_organization(self, directory: str) -> Dict[str, List[str]]:
        """
        Build a preview by making ONE LLM call for all files. 
        Then auto-create 'AutoLLM' rules if none exist for each extension/category.
        """
        # 1) Bulk classify
        filename_to_cat = self.bulk_classify_all_files(directory)
        # => { "file1.txt": "Business", "file2.pdf": "Academic/Math", ... }

        # 2) Build the preview
        preview = {}
        for fname, cat in filename_to_cat.items():
            preview.setdefault(cat, []).append(fname)

        # 3) AUTO-CREATE rules for each (extension, category) if needed
        for cat, files in preview.items():
            for f in files:
                ext = os.path.splitext(f)[1].lower()
                existing_rule = next(
                    (r for r in self.context.rules 
                     if r.name.startswith("AutoLLM_")
                        and r.category == cat
                        and (r.extension or '').lower() == ext),
                    None
                )
                if not existing_rule:
                    # Create an auto rule
                    rule_name = f"AutoLLM_{os.path.splitext(f)[0]}"
                    new_rule = Rule(
                        name=rule_name,
                        category=cat,
                        extension=ext,
                        keywords=[],
                        priority=1
                    )
                    self.context.add_rule(new_rule)
                    logger.info(f"Auto-saved new rule: {rule_name} => {cat} (ext={ext})")
                    if cat not in self.context.default_categories:
                        self.context.default_categories.append(cat)

        return preview

    # --------------------
    # The rest of the logic
    # --------------------

    def refine_preview(self, preview: Dict[str, List[str]], user_input: str) -> Dict[str, List[str]]:
        flatten_structure = any(
            kw in user_input.lower() 
            for kw in ["simple", "single", "no nested", "dont create subfolders", "dont use nested"]
        )
        move_request = re.search(r'move\s+(.+?)\s+to\s+(\w+)', user_input.lower())
        group_request = re.search(r'group\s+by\s+(\w+)', user_input.lower())

        new_preview = {}

        try:
            if flatten_structure or "move" in user_input.lower():
                assigned_files = set()
                file_categories = {}
                for cat_path, file_list in preview.items():
                    main_cat = cat_path.split('/')[0]  # flatten subfolders
                    for f in file_list:
                        file_categories.setdefault(f, set()).add(main_cat)

                if move_request:
                    file_pattern = move_request.group(1).strip()
                    target_cat = move_request.group(2).capitalize()
                    for filename in list(file_categories.keys()):
                        if file_pattern.lower() in filename.lower():
                            file_categories[filename] = {target_cat}

                for filename, cats in file_categories.items():
                    if len(cats) == 1:
                        final_cat = list(cats)[0]
                    else:
                        non_docs = [c for c in cats if c != 'Documents']
                        final_cat = non_docs[0] if non_docs else 'Documents'

                    if flatten_structure:
                        final_cat = final_cat.split('/')[0]

                    new_preview.setdefault(final_cat, [])
                    if filename not in assigned_files:
                        new_preview[final_cat].append(filename)
                        assigned_files.add(filename)

            elif group_request:
                group_by = group_request.group(1)
                if group_by in ['type', 'extension']:
                    for cat_path, file_list in preview.items():
                        for f in file_list:
                            ext = os.path.splitext(f)[1][1:].upper() or 'NO_EXTENSION'
                            group_cat = f"FileType_{ext}"
                            new_preview.setdefault(group_cat, []).append(f)

            if not new_preview:
                return preview

            cleaned_preview = {}
            for cat, flist in new_preview.items():
                clean_cat = cat.strip().replace(' ', '')
                if clean_cat not in cleaned_preview:
                    cleaned_preview[clean_cat] = []
                seen = set()
                for fname in flist:
                    if fname not in seen:
                        cleaned_preview[clean_cat].append(fname)
                        seen.add(fname)

            return {c: fs for c, fs in cleaned_preview.items() if fs}

        except Exception as e:
            logger.error(f"Error refining preview: {str(e)}")
            return preview

    def update_auto_rules(self, preview: Dict[str, List[str]]) -> None:
        extension_to_categories = {}
        for category, files in preview.items():
            for f in files:
                ext = os.path.splitext(f)[1].lower()
                extension_to_categories.setdefault(ext, []).append(category)

        for ext, cat_list in extension_to_categories.items():
            cat_count = {}
            for c in cat_list:
                cat_count[c] = cat_count.get(c, 0) + 1
            final_cat = sorted(cat_count.items(), key=lambda x: x[1], reverse=True)[0][0]

            for rule in self.context.rules:
                if rule.name.startswith("AutoLLM_") and (rule.extension or '').lower() == ext:
                    if rule.category != final_cat:
                        logger.info(f"Updating auto-rule '{rule.name}' from {rule.category} -> {final_cat}")
                        rule.category = final_cat
        self.context.rules.sort(key=lambda r: r.priority, reverse=True)

    def display_preview(self, preview: Dict[str, List[str]]) -> None:
        print("\nProposed File Organization:")
        print("==========================")
        for category in sorted(preview.keys()):
            print(f"\n📁 {category}")
            for filename in sorted(preview[category]):
                print(f"  📄 {filename}")
        print("\nOptions:")
        print("1. Proceed with sorting")
        print("2. Refine organization")
        print("3. Cancel")

    def organize_files(self, directory: str, preview: Dict[str, List[str]]) -> None:
        try:
            moves = {}
            for category, files in preview.items():
                category_path = os.path.join(directory, category)
                os.makedirs(category_path, exist_ok=True)

                for filename in files:
                    original_path = None
                    for root, _, filenames in os.walk(directory):
                        if filename in filenames:
                            original_path = os.path.join(root, filename)
                            break
                    if original_path and os.path.exists(original_path):
                        new_path = os.path.join(category_path, filename)
                        moves[original_path] = new_path

            for src, dst in moves.items():
                if os.path.exists(src):
                    old_rel_dir = os.path.relpath(os.path.dirname(src), directory)
                    fname = os.path.basename(src)
                    if old_rel_dir in self.original_files and fname in self.original_files[old_rel_dir]:
                        self.original_files[old_rel_dir].remove(fname)

                    os.makedirs(os.path.dirname(dst), exist_ok=True)
                    if os.path.exists(dst):
                        base, ext = os.path.splitext(dst)
                        counter = 1
                        while os.path.exists(f"{base}_{counter}{ext}"):
                            counter += 1
                        dst = f"{base}_{counter}{ext}"

                    shutil.move(src, dst)
                    logger.info(f"Moved {src} to {dst}")

                    new_rel_dir = os.path.relpath(os.path.dirname(dst), directory)
                    self.original_files.setdefault(new_rel_dir, []).append(os.path.basename(dst))
                else:
                    logger.warning(f"Source file not found: {src}")

            self.cleanup_empty_folders(directory)
            print("\n✅ Files have been organized successfully!")
        
        except Exception as e:
            logger.error(f"Error organizing files: {str(e)}")
            print("\n⚠️ Error occurred while organizing files. Please check the logs.")

    def cleanup_empty_folders(self, directory: str) -> None:
        try:
            for root, dirs, files in os.walk(directory, topdown=False):
                for d in dirs:
                    dir_path = os.path.join(root, d)
                    if not os.listdir(dir_path):
                        os.rmdir(dir_path)
                        logger.info(f"Removed empty directory: {dir_path}")
        except Exception as e:
            logger.error(f"Error during cleanup: {str(e)}")

    def revert_organization(self, directory: str) -> None:
        try:
            moved_files = []
            for root, _, files in os.walk(directory):
                if root != directory:
                    for f in files:
                        moved_files.append((os.path.join(root, f), f))

            for current_path, filename in moved_files:
                target_path = os.path.join(directory, filename)
                if os.path.exists(current_path):
                    if os.path.exists(target_path):
                        base, ext = os.path.splitext(filename)
                        counter = 1
                        while os.path.exists(os.path.join(directory, f"{base}_{counter}{ext}")):
                            counter += 1
                        target_path = os.path.join(directory, f"{base}_{counter}{ext}")

                    shutil.move(current_path, target_path)
                    logger.info(f"Reverted {current_path} to {target_path}")

                    old_rel_dir = os.path.relpath(os.path.dirname(current_path), directory)
                    if old_rel_dir in self.original_files and filename in self.original_files[old_rel_dir]:
                        self.original_files[old_rel_dir].remove(filename)

                    new_rel_dir = ""
                    self.original_files.setdefault(new_rel_dir, []).append(os.path.basename(target_path))

            for root, dirs, _ in os.walk(directory, topdown=False):
                for d in dirs:
                    dir_path = os.path.join(root, d)
                    try:
                        os.rmdir(dir_path)
                        logger.info(f"Removed empty directory: {dir_path}")
                    except OSError:
                        pass

            self.cleanup_empty_folders(directory)
            print("\n✅ Successfully reverted file organization!")
        
        except Exception as e:
            logger.error(f"Error reverting organization: {str(e)}")
            print("\n⚠️ Error occurred while reverting files. Please check the logs.")

    def show_rules(self) -> None:
        if not self.context.rules:
            print("\n❌ No active organization rules.")
            return
        print("\n📋 Current Organization Rules:")
        print("=" * 40)
        for i, rule in enumerate(self.context.rules, 1):
            print(f"📌 Rule {i}:")
            print(f"   🏷️  Name: {rule.name}")
            print(f"   📁 Category: {rule.category}")
            if rule.extension:
                print(f"   📎 Extension: {rule.extension}")
            if rule.date_range:
                print(f"   📅 Date Range: {rule.date_range}")
            if rule.keywords:
                print(f"   🔍 Keywords: {rule.keywords}")
            print(f"   ⭐ Priority: {rule.priority}")
            print("-" * 40)

    def revert_rule(self, rule_index: int, directory: str) -> Optional[Dict[str, List[str]]]:
        try:
            if 0 <= rule_index < len(self.context.rules):
                removed_rule = self.context.rules.pop(rule_index)
                print(f"\nReverted rule: {removed_rule.name}")
                # Return updated preview
                return self.preview_organization(directory)
            else:
                print("\nInvalid rule index.")
                return None
        except Exception as e:
            logger.error(f"Error reverting rule: {str(e)}")
            return None

    def run(self, directory: str) -> None:
        """
        Same menu flow, but preview_organization() 
        does a single bulk classification for all files.
        """
        try:
            self.original_files.clear()
            for root, _, files in os.walk(directory):
                rel_path = os.path.relpath(root, directory)
                if rel_path == '.':
                    rel_path = ''
                self.original_files[rel_path] = [f for f in files]

            while True:
                print("\n" + "="*50)
                print("🗂️  File Organizer Menu".center(50))
                print("="*50)
                print("1. 📁 Organize files")
                print("2. 🔄 Revert organization")
                print("3. ⚙️  Manage rules")
                print("4. 🚪 Exit")
                print("-"*50)

                choice = input("\nEnter your choice (1-4): ")
                if choice == "1":
                    preview = self.preview_organization(directory)
                    while True:
                        self.display_preview(preview)
                        print("\n" + "-"*50)
                        print("📋 Organization Options:")
                        print("1. ✅ Proceed with sorting")
                        print("2. 🔄 Refine organization")
                        print("3. ❌ Cancel")
                        print("-"*50)

                        subchoice = input("\nEnter your choice (1-3): ")
                        if subchoice == "1":
                            self.organize_files(directory, preview)
                            break
                        elif subchoice == "2":
                            user_input = input("\nHow would you like to modify the organization? ")
                            preview = self.refine_preview(preview, user_input)
                            self.update_auto_rules(preview)
                        elif subchoice == "3":
                            print("\n❌ Operation cancelled.")
                            break
                        else:
                            print("\n⚠️ Invalid choice. Try again.")

                elif choice == "2":
                    confirm = input("\n⚠️  Revert all files to the root folder? (y/n): ")
                    if confirm.lower() == 'y':
                        self.revert_organization(directory)
                    else:
                        print("\n❌ Revert cancelled.")

                elif choice == "3":
                    while True:
                        print("\n" + "-"*50)
                        print("⚙️  Rule Management:")
                        print("1. 👀 View current rules")
                        print("2. ➕ Add new rule")
                        print("3. ❌ Remove a rule")
                        print("4. 🔙 Back to main menu")
                        print("-"*50)

                        rule_choice = input("\nEnter your choice (1-4): ")
                        if rule_choice == "1":
                            self.show_rules()
                        elif rule_choice == "2":
                            rule_text = input("\nDescribe your rule: ")
                            success = self.add_custom_rule(rule_text)
                            if success:
                                print("✅ Rule added successfully!")
                            else:
                                print("❌ Failed to add rule.")
                        elif rule_choice == "3":
                            self.show_rules()
                            if self.context.rules:
                                try:
                                    idx = int(input("\nEnter the rule number to remove (0=cancel): ")) - 1
                                    if 0 <= idx < len(self.context.rules):
                                        removed_rule = self.context.rules.pop(idx)
                                        print(f"\n✅ Removed rule: {removed_rule.name}")
                                    else:
                                        print("\n❌ Cancelled.")
                                except ValueError:
                                    print("\n⚠️ Invalid input.")
                        elif rule_choice == "4":
                            break
                        else:
                            print("\n⚠️ Invalid choice. Try again.")

                elif choice == "4":
                    print("\n👋 Exiting File Organizer.")
                    return
                else:
                    print("\n⚠️ Invalid choice. Try again.")

        except Exception as e:
            logger.error(f"Error in file organizer: {str(e)}")
            print("\n⚠️  An error occurred. Check logs for details.")
